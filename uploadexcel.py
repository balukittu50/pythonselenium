import openpyxl
from selenium import webdriver
from selenium.webdriver.common.by import By
from selenium.webdriver.support import expected_conditions
from selenium.webdriver.support.wait import WebDriverWait


def updateexcel(filepath, searchterm, colname, newvalue):
    book = openpyxl.load_workbook(filepath)
    sheet = book.active
    Dic = {}

    for i in range(1, sheet.max_column + 1):
        if sheet.cell(row=1, column=i).value == colname:
            Dic["col"] = i

    for i in range(1, sheet.max_row + 1):
        for j in range(1, sheet.max_column + 1):
            if sheet.cell(row=i, column=j).value == searchterm:
                Dic["row"] = i

    sheet.cell(row=Dic["row"], column=Dic["col"]).value = newvalue
    book.save(filepath)


filepath = "F:\\COURSES\\Selenium java\\UdemySelinumPython\\download.xlsx"
fruitname = "Apple"
newvalue = 123
driver = webdriver.Chrome()
driver.get("https://rahulshettyacademy.com/upload-download-test/index.html")
driver.maximize_window()
driver.implicitly_wait(5)

driver.find_element(By.ID, "downloadButton").click()

# edit the excel updated value
updateexcel(filepath, fruitname, "price", newvalue)

# upload data
file_input = driver.find_element(By.CSS_SELECTOR, "input[type='file']")
file_input.send_keys(filepath)

wait = WebDriverWait(driver, 5)
toastlocator = (By.CSS_SELECTOR, ".Toastify__toast-body div:nth-child(2)")
wait.until(expected_conditions.visibility_of_element_located(toastlocator))

print(driver.find_element(*toastlocator).text)
price_colmn = driver.find_element(By.XPATH, "//div[text()='Price']").get_attribute("data-column-id")
actual_price = driver.find_element(By.XPATH,
                                   "//div[text()='" + fruitname + "']/parent::div/parent::div/div[@id='cell-" + price_colmn + "-undefined']").text
print(actual_price)

assert actual_price == str(newvalue)
